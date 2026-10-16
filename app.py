# Note: This file merges app.py (Diagnostic & Interventional Radiology incl. Nuclear Medicine) 
# and app_rt.py (Radiotherapy Physics) into a single unified Streamlit application. 
# Shared tables (A & B) are defined once, while C1 and C2 are kept strictly separated 
# with independent per-specialty answer stores and dynamic Section 4b report generators.
#testing
import streamlit as st
import json
import pandas as pd
import io
from docx import Document
from datetime import datetime

# ==============================================================================
# KSC DATASETS (SHARED TABLES A & B, SEPARATE TABLES C1 & C2)
# ==============================================================================

TABLE_A = [
    {
        "table": "Table A", "code": "A1.1", "title": "Radiation Physics & Matter Interactions",
        "description": "Comprehensive understanding of atomic/nuclear physics, radiation production, photon/electron/hadron interactions with matter, and attenuation physics in tissue."
    },
    {
        "table": "Table A", "code": "A1.2", "title": "Radiation Dosimetry Principles & Cavity Theory",
        "description": "In-depth knowledge of cavity theory (Spencer-Attix, Bragg-Gray), absolute and relative dosimetry instrumentation (ionization chambers, diodes, MOSFETs, film, gels)."
    },
    {
        "table": "Table A", "code": "A1.3", "title": "Radiation Detection & Measurement Instrumentation",
        "description": "Operating principles, calibration techniques, energy dependence, directional response, and signal processing of radiation detection systems."
    },
    {
        "table": "Table A", "code": "A2.1", "title": "Legislative & Regulatory Frameworks",
        "description": "Detailed mastery of Irish radiation protection legislation (S.I. 256 of 2018, S.I. 30 of 2019), EU Directives (2013/59/Euratom), and HIQA framework compliance."
    },
    {
        "table": "Table A", "code": "A2.2", "title": "Radiation Safety Administration & Shielding Design",
        "description": "Design and calculation of structural radiation shielding for high-energy linacs (NCRP 151) and brachytherapy suites; safety committee governance and hazard management."
    },
    {
        "table": "Table A", "code": "A3.1", "title": "Anatomy, Physiology & Medical Imaging Physics",
        "description": "Applied knowledge of cross-sectional human anatomy, radio-pathology, and medical imaging modalities (CT, MRI, PET, SPECT) used in clinical planning and diagnostics."
    },
    {
        "table": "Table A", "code": "A3.2", "title": "Biological Effects of Radiation & Radiobiology",
        "description": "Linear-Quadratic model physics, 5 Rs of radiobiology, cell survival curves, fractionations, TCP/NTCP modeling, and acute/late tissue toxicity mechanisms."
    },
    {
        "table": "Table A", "code": "A4.1", "title": "Medical Informatics, Networking & Cybersecurity",
        "description": "DICOM/DICOM-RT standards, PACS/OIS database management, system networking, data security, HIPAA/GDPR compliance, and algorithm software validation."
    }
]

TABLE_B = [
    {
        "table": "Table B", "code": "B1.1", "title": "Clinical Governance & Quality Management Systems",
        "description": "Implementation of comprehensive departmental QMS, ISO quality standards, clinical audit protocols, and document control procedures."
    },
    {
        "table": "Table B", "code": "B1.2", "title": "Risk Management & Patient Safety",
        "description": "Proactive risk assessment methodologies (FMEA, FTA), incident reporting systems, root cause analysis (RCA), and human factors engineering."
    },
    {
        "table": "Table B", "code": "B2.1", "title": "Ethics, Patient Autonomy & Bioethics",
        "description": "Adherence to medical ethics, patient confidentiality, informed consent frameworks in medical exposures, and professional code of conduct."
    },
    {
        "table": "Table B", "code": "B2.2", "title": "Communication, Leadership & Interdisciplinary Collaboration",
        "description": "Effective technical and clinical communication with clinicians, technologists, nursing staff, hospital management, and regulatory authorities."
    },
    {
        "table": "Table B", "code": "B3.1", "title": "Research, Innovation & Health Technology Assessment",
        "description": "Design of clinical research protocols, statistical analysis, peer-reviewed scientific publishing, evaluation of emerging technologies, and cost-benefit analysis."
    },
    {
        "table": "Table B", "code": "B3.2", "title": "Continuing Professional Development (CPD) & Education",
        "description": "Maintaining personal CPD portfolios, delivering technical training to multidisciplinary teams, and mentoring junior staff and trainees."
    }
]

TABLE_C1 = [
    {
        "table": "Table C1", "code": "C1.1", "title": "Diagnostic Scientific Problem Solving Service",
        "description": "Statutory requirements, ionising radiation properties, radiobiological models, image quality measurement, inverse problems, and image reconstruction algorithms."
    },
    {
        "table": "Table C1", "code": "C1.2", "title": "Diagnostic Dosimetry Measurements",
        "description": "Patient dosimetric quantities, operational/public dose indicators, MIRD formalism, tumour dosimetry, and dosimeter calibration."
    },
    {
        "table": "Table C1", "code": "C1.3", "title": "Diagnostic Patient Safety / Dose Optimisation",
        "description": "Radiobiological dose-effect relationships, justification and optimization principles, diagnostic reference levels (DRLs), high-dose practices, and incident investigations."
    },
    {
        "table": "Table C1", "code": "C1.4", "title": "Clinical Device Management in Diagnostic Radiology",
        "description": "Technological infrastructure, acceptance testing, commissioning, quality control, ICT/PACS integration, and decontamination in Nuclear Medicine."
    },
    {
        "table": "Table C1", "code": "C1.5", "title": "Clinical Involvement in Diagnostic Radiology",
        "description": "Anatomical/functional imaging modalities, acquisition protocol manipulation, clinical applications, image post-processing, and multi-disciplinary teamwork."
    },
    {
        "table": "Table C1", "code": "C1.6", "title": "Development of Service Quality and Cost-Effectiveness",
        "description": "Service quality development, quality management performance, cross-institutional quality control, and evidence-based medicine."
    },
    {
        "table": "Table C1", "code": "C1.7", "title": "Expert Consultancy in Diagnostic Radiology",
        "description": "Ethical issues in expert consultancy, role of scientists as consultants, and independent report auditing."
    },
    {
        "table": "Table C1", "code": "C1.8", "title": "Education of Healthcare Professionals in Diagnostic Radiology",
        "description": "Adult pedagogy, staff training programs in radiation safety, and knowledge transfer for healthcare professionals and trainees."
    },
    {
        "table": "Table C1", "code": "C1.9", "title": "Health Technology Assessment in Diagnostic Radiology",
        "description": "Systematic reviews, clinical trials administration, research medical exposures, and statistical analysis for HTA."
    },
    {
        "table": "Table C1", "code": "C1.10", "title": "Innovation in Diagnostic Radiology",
        "description": "Horizon scanning for new/emerging technologies, device integration, software modification, and new experimental set-ups."
    }
]

TABLE_C2 = [
    {
        "table": "Table C2", "code": "C2.1", "title": "Scientific Problem Solving Service",
        "description": "Explains statutory requirements, operation of treatment devices (cobalt units, medical linacs, robotic linacs, particle systems), external beam techniques, in-room imaging, and algorithm limitations."
    },
    {
        "table": "Table C2", "code": "C2.2", "title": "External Beam & Brachytherapy Dosimetry Measurement",
        "description": "Covers dosimetry terminology (PDD, TMR, TPR), IAEA absorbed dose protocols, in-vivo dosimetry, calibration chains, reference dosimetry, and source strength determinations."
    },
    {
        "table": "Table C2", "code": "C2.3", "title": "Patient Safety / Dose Optimisation",
        "description": "Explains radiobiological dose-effect relationships, tissue responses, treatment planning, and optimization principles for photons, electrons, protons, heavy ions, and brachytherapy."
    },
    {
        "table": "Table C2", "code": "C2.4", "title": "Occupational & Public Safety / Dose Optimisation",
        "description": "Covers statutory requirements, occupational/public risk estimation in radiotherapy, source transport/handling regulations, and safety plan design."
    },
    {
        "table": "Table C2", "code": "C2.5", "title": "Clinical Medical Device Management in Radiation Oncology",
        "description": "Explains hardware/software specifications for treatment units, TPS, networking standards (DICOM/DICOM-RT), acceptance testing, commissioning, quality control, and source management."
    },
    {
        "table": "Table C2", "code": "C2.6", "title": "Clinical Involvement in Radiation Oncology",
        "description": "Covers oncogenesis, radiobiology models, TCP/NTCP, ICRU volume definitions (GTV, CTV, PTV), plan optimization, IMRT/VMAT/SBRT delivery, motion management, and brachytherapy."
    },
    {
        "table": "Table C2", "code": "C2.7", "title": "Development of Service Quality & Cost-Effectiveness in RO",
        "description": "Covers statutory requirements and leadership in multi-disciplinary quality development across all steps of the radiotherapy treatment chain (simulation, planning, verification, delivery, reporting)."
    },
    {
        "table": "Table C2", "code": "C2.8", "title": "Expert Consultancy in Radiation Oncology",
        "description": "Focuses on statutory requirements and management of professional and ethical issues associated with consultancy involving high levels of ionizing radiation."
    },
    {
        "table": "Table C2", "code": "C2.9", "title": "Education of Healthcare Professionals in RO",
        "description": "Covers statutory requirements and specialized education/training addressing the clinical use of high levels of ionizing radiation for healthcare professionals and trainees."
    },
    {
        "table": "Table C2", "code": "C2.10", "title": "Technology Assessment in Radiation Oncology",
        "description": "Covers statutory requirements, health technology assessments (HTA), and management of research medical exposures and clinical trials involving high levels of ionizing radiation."
    },
    {
        "table": "Table C2", "code": "C2.11", "title": "Innovation in Radiation Oncology",
        "description": "Focuses on identification, management, and responsibility for legal and technical issues associated with innovation and new developments in radiation oncology and high-dose radiation."
    }
]

RATING_OPTIONS = [
    "5 - Master: Departmental lead, innovator, national subject authority",
    "4 - Expert: Independent clinical expert, routine performer, supervisor level",
    "3 - Proficient: Competent execution under general supervision",
    "2 - Foundation: Theoretical knowledge, limited clinical practice",
    "1 - In Development: Gap identified; active clinical training required"
]

SPECIALTIES_CONFIG = {
    "Diagnostic & Interventional Radiology": {
        "display_name": "Diagnostic & Interventional Radiology (incl. Nuclear Medicine)",
        "domain_label": "Diagnostic & Interventional Radiology",
        "c_table": TABLE_C1,
        "matrix_caption": "Tables A, B & C1 Summary"
    },
    "Radiotherapy Physics": {
        "display_name": "Radiotherapy Physics (Radiation Oncology)",
        "domain_label": "Radiotherapy Physics",
        "c_table": TABLE_C2,
        "matrix_caption": "Tables A, B & C2 Summary"
    }
}

# ==============================================================================
# APPLICATION STATE & HELPER FUNCTIONS
# ==============================================================================

def init_state():
    """Initializes Streamlit session state with per-specialty answer stores."""
    if "candidate_name" not in st.session_state: st.session_state.candidate_name = ""
    if "hospital" not in st.session_state: st.session_state.hospital = ""
    if "supervisors" not in st.session_state: st.session_state.supervisors = ""
    if "selected_specialty" not in st.session_state: 
        st.session_state.selected_specialty = "Diagnostic & Interventional Radiology"
    if "answers" not in st.session_state: 
        st.session_state.answers = {
            "Diagnostic & Interventional Radiology": {},
            "Radiotherapy Physics": {}
        }
    if "current_idx" not in st.session_state: st.session_state.current_idx = 0

def get_active_dataset():
    """Returns the combined list of requirements for the currently selected specialty."""
    spec_key = st.session_state.selected_specialty
    c_table = SPECIALTIES_CONFIG[spec_key]["c_table"]
    return TABLE_A + TABLE_B + c_table

def reset_state():
    """Clears the session state for the active specialty."""
    spec = st.session_state.selected_specialty
    st.session_state.candidate_name = ""
    st.session_state.hospital = ""
    st.session_state.supervisors = ""
    st.session_state.answers[spec] = {}
    st.session_state.current_idx = 0

def handle_specialty_change():
    """Callback function triggered instantly when the selectbox value changes."""
    st.session_state.current_idx = 0

def handle_backup_upload():
    """Callback function triggered when a JSON backup is uploaded."""
    uploaded_file = st.session_state.get("backup_uploader")
    if uploaded_file is not None:
        try:
            data = json.load(uploaded_file)
            restored_spec = data.get("specialty")
            if restored_spec in SPECIALTIES_CONFIG:
                st.session_state.selected_specialty = restored_spec
            
            st.session_state.candidate_name = data.get("candidate_name", "")
            st.session_state.hospital = data.get("hospital", "")
            st.session_state.supervisors = data.get("supervisors", "")
            
            restored_answers = data.get("answers", {})
            if isinstance(restored_answers, dict):
                if any(k.startswith("A") or k.startswith("B") or k.startswith("C") for k in restored_answers.keys()):
                    st.session_state.answers[st.session_state.selected_specialty] = restored_answers
                else:
                    st.session_state.answers = restored_answers
                    
            st.session_state.current_idx = min(
                data.get("current_idx", 0), 
                len(get_active_dataset()) - 1
            )
            st.toast("Progress restored successfully!", icon="✅")
        except Exception as e:
            st.error(f"Failed to parse JSON backup file: {e}")

def generate_word_report():
    """Generates a Word document including the competency matrix table and detailed audit data."""
    doc = Document()
    today_str = datetime.now().strftime("%Y-%m-%d")
    spec_key = st.session_state.selected_specialty
    domain_label = SPECIALTIES_CONFIG[spec_key]["domain_label"]
    active_data = get_active_dataset()
    current_answers = st.session_state.answers.get(spec_key, {})
    
    doc.add_heading("MPE Equivalence Demonstration Audit Report", level=0)
    
    doc.add_paragraph(f"Specialty Domain:  {domain_label}")
    doc.add_paragraph(f"Generated Date: {today_str}")
    
    doc.add_heading("1. Candidate & Clinical Infrastructure", level=1)
    doc.add_paragraph(f"Candidate Name: {st.session_state.candidate_name or 'N/A'}", style="List Bullet")
    doc.add_paragraph(f"Hospital / Institution: {st.session_state.hospital or 'N/A'}", style="List Bullet")
    doc.add_paragraph(f"Supervising MPE Supporter(s): {st.session_state.supervisors or 'N/A'}", style="List Bullet")
    
    # Competency Overview Matrix Table Summary included in Word report
    doc.add_heading("2. Competency Overview Matrix Summary", level=1)
    matrix_table = doc.add_table(rows=1, cols=3)
    matrix_table.style = 'Table Grid'
    hdr_cells = matrix_table.rows[0].cells
    hdr_cells[0].text = "Code"
    hdr_cells[1].text = "Requirement Title"
    hdr_cells[2].text = "Self-Rating"
    
    for k in active_data:
        ans = current_answers.get(k["code"], {})
        rating_val = ans.get("rating", "Unrated")
        row_cells = matrix_table.add_row().cells
        row_cells[0].text = k["code"]
        row_cells[1].text = k["title"]
        row_cells[2].text = rating_val
        
    doc.add_paragraph() # Spacing
    
    doc.add_heading("3. Detailed Knowledge, Skills & Competency (KSC) Audit Matrix", level=1)
    
    for k in active_data:
        ans = current_answers.get(k["code"], {})
        rating = ans.get("rating", "Unrated / Pending Evaluation")
        evidence = ans.get("evidence", "No clinical documentary evidence provided.")
        
        doc.add_heading(f"[{k['code']}] {k['title']}", level=2)
        doc.add_paragraph(f"Table Category: {k['table']}")
        doc.add_paragraph(f"Requirement Scope: {k['description']}")
        doc.add_paragraph(f"Assessed Competency Level: {rating}")
        doc.add_paragraph("Clinical Documentary Evidence:")
        doc.add_paragraph(evidence, style="Quote")

    doc.add_heading("4. Applicant & Supervisor Declarations", level=1)
    doc.add_paragraph("I confirm that the ratings and clinical documentary evidence detailed in this cross-reference report accurately reflect my professional experience and practice.\n")
    
    doc.add_paragraph(f"Candidate Signature: ___________________________    Date: {today_str}")
    doc.add_paragraph("Supervising MPE Signature: ______________________    Date: _______________")
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ==============================================================================
# MAIN APPLICATION LAYOUT & SIDEBAR CONFIGURATION
# ==============================================================================

st.set_page_config(page_title="MPE Competency Assessment Suite", layout="wide")
init_state()

active_ksc_data = get_active_dataset()
current_answers = st.session_state.answers.setdefault(st.session_state.selected_specialty, {})

# Calculate Progress for Active Specialty
completed_count = sum(1 for k in active_ksc_data if k["code"] in current_answers and current_answers[k["code"]].get("rating"))
total_count = len(active_ksc_data)
progress_pct = completed_count / total_count if total_count > 0 else 0

# Sidebar Configuration
with st.sidebar:
    st.header("Applicant Profile")
    st.session_state.candidate_name = st.text_input("Candidate Name", value=st.session_state.candidate_name, placeholder="Dr. Jane Doe")
    st.session_state.hospital = st.text_input("Hospital / Institution", value=st.session_state.hospital, placeholder="Galway University Hospital")
    st.session_state.supervisors = st.text_input("Supervising MPE Supporter(s)", value=st.session_state.supervisors, placeholder="Prof. J. Smith, MPE")
    
    st.divider()
    st.subheader("Specialty Framework")
    specialty_options = list(SPECIALTIES_CONFIG.keys())
    
    st.selectbox(
        "Select MPE Specialty Framework",
        options=specialty_options,
        format_func=lambda x: SPECIALTIES_CONFIG[x]["display_name"],
        key="selected_specialty",
        on_change=handle_specialty_change,
        label_visibility="collapsed"
    )

    st.divider()
    st.metric(label="Total Completion", value=f"{int(progress_pct * 100)}%", delta=f"{completed_count} of {total_count} rated", delta_color="normal")
    st.progress(progress_pct)
    
    st.divider()
    st.subheader("Save & Restore Progress")
    
    export_data = json.dumps({
        "specialty": st.session_state.selected_specialty,
        "candidate_name": st.session_state.candidate_name,
        "hospital": st.session_state.hospital,
        "supervisors": st.session_state.supervisors,
        "answers": st.session_state.answers,
        "current_idx": st.session_state.current_idx
    }, indent=2)
    
    st.download_button(
        "Download Backup JSON", 
        data=export_data, 
        file_name=f"MPE_Backup_{st.session_state.selected_specialty.replace(' ', '_')}.json", 
        mime="application/json",
        use_container_width=True
    )
    
    st.file_uploader(
        "Restore from Backup", 
        type=["json"], 
        key="backup_uploader", 
        on_change=handle_backup_upload
    )
            
    st.divider()
    if st.button("Reset All Data", type="secondary", use_container_width=True):
        reset_state()
        st.rerun()

# Native Tab Navigation
tab1, tab2, tab3 = st.tabs(["Sequential Evaluation", "Competency Matrix", "Audit Report"])

# ------------------------------------------------------------------------------
# TAB 1: SEQUENTIAL EVALUATION
# ------------------------------------------------------------------------------
with tab1:
    col_jump, col_blank = st.columns([2, 1])
    with col_jump:
        jump_options = [f"[{k['code']}] {k['title']}" for k in active_ksc_data]
        selected_jump = st.selectbox(
            "Jump to Requirement", 
            options=range(len(jump_options)), 
            format_func=lambda x: jump_options[x], 
            index=min(st.session_state.current_idx, len(jump_options) - 1)
        )
        if selected_jump != st.session_state.current_idx:
            st.session_state.current_idx = selected_jump
            st.rerun()

    current_ksc = active_ksc_data[st.session_state.current_idx]
    current_code = current_ksc["code"]
    current_ans = current_answers.get(current_code, {"rating": None, "evidence": ""})

    col1, col2, col3 = st.columns(3)
    with col1:
        st.markdown("**Specialty Domain**")
        st.markdown(f"### {SPECIALTIES_CONFIG[st.session_state.selected_specialty]['domain_label']}")
    with col2:
        st.markdown("**Current Requirement**")
        st.markdown(f"### {current_code} ({st.session_state.current_idx + 1} of {len(active_ksc_data)})")
    with col3:
        st.markdown("**Table Section**")
        st.markdown(f"### {current_ksc['table']}")

    st.subheader(f"[{current_code}] {current_ksc['title']}")
    st.info(current_ksc["description"])

    def save_evaluation():
        rating_val = st.session_state.get(f"f_rating_{current_code}")
        evidence_val = st.session_state.get(f"f_evidence_{current_code}")
        
        current_answers[current_code] = {
            "rating": rating_val.split(":")[0] if rating_val else None,
            "evidence": evidence_val
        }
    
    default_rating_idx = 0
    if current_ans["rating"]:
        for i, opt in enumerate(RATING_OPTIONS):
            if current_ans["rating"] in opt:
                default_rating_idx = i
                break

    st.radio(
        "1. Self-Assessed Competency Level", 
        options=RATING_OPTIONS, 
        index=default_rating_idx, 
        key=f"f_rating_{current_code}", 
        on_change=save_evaluation
    )
    
    st.text_area(
        "2. Clinical Documentary Evidence", 
        value=current_ans["evidence"], 
        height=150, 
        key=f"f_evidence_{current_code}", 
        on_change=save_evaluation, 
        placeholder="Detail specific clinical tasks, dosimetry modeling, quality control protocols, or published reports..."
    )

    # Navigation Controls
    c1, c2, c3 = st.columns([1, 8, 1])
    
    with c1:
        if st.session_state.current_idx > 0:
            if st.button("Previous", use_container_width=True):
                st.session_state.current_idx -= 1
                st.rerun()
                
    with c3:
        if st.session_state.current_idx < len(active_ksc_data) - 1:
            if st.button("Next Item", type="primary", use_container_width=True):
                st.session_state.current_idx += 1
                st.rerun()

# ------------------------------------------------------------------------------
# TAB 2: COMPETENCY MATRIX
# ------------------------------------------------------------------------------
with tab2:
    st.subheader("Competency Overview Matrix")
    st.caption(SPECIALTIES_CONFIG[st.session_state.selected_specialty]["matrix_caption"])
    
    matrix_data = []
    for k in active_ksc_data:
        ans = current_answers.get(k["code"], {})
        matrix_data.append({
            "Code": k["code"],
            "Requirement Title": k["title"],
            "Self-Rating": ans.get("rating", "Unrated"),
            "Clinical Evidence": ans.get("evidence", "No evidence recorded.")
        })
    
    df = pd.DataFrame(matrix_data)
    st.dataframe(df, use_container_width=True, hide_index=True)

# ------------------------------------------------------------------------------
# TAB 3: AUDIT REPORT
# ------------------------------------------------------------------------------
with tab3:
    st.subheader("Audit Summary Report")
    
    word_buffer = generate_word_report()
    
    st.download_button(
        label="Download Word Audit File",
        data=word_buffer,
        file_name=f"MPE_Section4b_Audit_Report_{st.session_state.candidate_name.replace(' ', '_') or 'Draft'}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )

    st.markdown("*(The report content is compiled directly into the Word document available for download above.)*")
